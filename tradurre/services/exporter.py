from io import BytesIO

from docx import Document

from tradurre.services.html_utils import strip_html


def export_docx(pairs: list[dict], mode: str) -> bytes:
    doc = Document()
    for pair in pairs:
        if mode == "source":
            doc.add_paragraph(strip_html(pair["source_html"]))
        elif mode == "target":
            doc.add_paragraph(strip_html(pair["target_html"]))
        elif mode == "parallel":
            doc.add_paragraph(strip_html(pair["source_html"]))
            doc.add_paragraph(strip_html(pair["target_html"]))
            doc.add_paragraph("")  # blank separator

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_txt(pairs: list[dict], mode: str) -> str:
    blocks = []
    for pair in pairs:
        if mode == "source":
            blocks.append(strip_html(pair["source_html"]))
        elif mode == "target":
            blocks.append(strip_html(pair["target_html"]))
        elif mode == "parallel":
            src = strip_html(pair["source_html"])
            tgt = strip_html(pair["target_html"])
            blocks.append(f"{src}\n---\n{tgt}")
    return "\n\n".join(blocks)
