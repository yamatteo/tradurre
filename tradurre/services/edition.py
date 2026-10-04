"""Edition export (SPEC §3.5; PLAN.md, "Edition export"): one side's text as reviewed, as .txt or .docx.

The edition is the current `text` of every segment (edits included, never `original_text`) of the blocks that are not
excluded, in document order, one paragraph per block.
"""

import io
import sqlite3

from docx import Document


def edition_blocks(conn: sqlite3.Connection, book_id: str, side: str) -> list[tuple[str, str]]:
    """(block kind, the block's segment texts joined by one space) for every included block of `side`, in order."""
    rows = conn.execute(
        "SELECT b.id, b.kind, s.text FROM blocks b "
        "JOIN documents d ON d.id = b.document_id "
        "JOIN segments s ON s.block_id = b.id "
        "WHERE d.project_id = ? AND d.side = ? AND b.excluded = 0 "
        "ORDER BY b.ord, s.ord",
        (book_id, side),
    ).fetchall()
    blocks: list[tuple[int, str, list[str]]] = []
    for row in rows:
        if not blocks or blocks[-1][0] != row["id"]:
            blocks.append((row["id"], row["kind"], []))
        blocks[-1][2].append(row["text"])
    return [(kind, " ".join(texts)) for _, kind, texts in blocks]


def edition_txt(blocks: list[tuple[str, str]]) -> bytes:
    """UTF-8 with a BOM (so Notepad and Word on Windows read the encoding right), a blank line between blocks."""
    return ("﻿" + "\n\n".join(text for _, text in blocks) + "\n").encode("utf-8")


def edition_docx(blocks: list[tuple[str, str]]) -> bytes:
    """A heading block becomes a level-1 heading, any other a paragraph."""
    doc = Document()
    for kind, text in blocks:
        if kind == "heading":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
