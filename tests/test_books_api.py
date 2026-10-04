"""Tests for the books API (/api/v2): import into the new model and read back."""

from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from tradurre.app import app
from tradurre.db import get_connection
from tradurre.domain.blocks import exclude_block
from tradurre.domain.history import transaction

SOURCE = "Marie arriva.\n\nIl pleuvait.\n\nPaul partit. Il ne dit rien."
TARGET = "Marie arrivò.\n\nPaul partì. Non disse niente."


@pytest.fixture()
def db_path(tmp_path, monkeypatch):
    path = tmp_path / "test.db"
    monkeypatch.setattr("tradurre.config.DB_PATH", path)
    monkeypatch.setattr("tradurre.app.DB_PATH", path)
    return path


@pytest.fixture()
def client(db_path):
    with TestClient(app) as c:
        yield c


def _import(client, source=("contrefeu.txt", SOURCE.encode()), target=("sacro.txt", TARGET.encode()), **form):
    return client.post("/api/v2/books", files={"source": source, "target": target}, data=form)


def _texts(bead, side):
    return [s["text"] for s in bead[side]]


def _docx(*paragraphs):
    doc = Document()
    for kind, text in paragraphs:
        if kind == "heading":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_import_txt_and_read_back(client):
    resp = _import(client, title="Contrefeu", source_lang="fr", target_lang="it")
    assert resp.status_code == 201
    body = resp.json()
    assert (body["title"], body["bead_count"], body["warnings"]) == ("Contrefeu", 3, [])

    book = client.get(f"/api/v2/books/{body['id']}").json()
    assert (book["title"], book["source_lang"], book["target_lang"]) == ("Contrefeu", "fr", "it")
    assert [(_texts(b, "source"), _texts(b, "target")) for b in book["beads"]] == [
        # Too little text for the length aligner to leave "Il pleuvait." unpaired (see test_build.py).
        (["Marie arriva."], ["Marie arrivò."]),
        (["Il pleuvait."], ["Paul partì."]),
        (["Paul partit.", "Il ne dit rien."], ["Non disse niente."]),
    ]
    first = book["beads"][0]
    assert (first["method"], first["reviewed"]) == ("length", False)
    assert 0.8 < first["confidence"] <= 1
    assert first["source"][0]["block_kind"] == "paragraph"
    assert book["excluded"] == []


def test_import_docx_with_heading(client):
    resp = _import(
        client,
        source=("livre.docx", _docx(("heading", "Chapitre premier"), ("paragraph", "Marie arriva."))),
        target=("libro.docx", _docx(("heading", "Capitolo primo"), ("paragraph", "Marie arrivò."))),
    )
    assert resp.status_code == 201
    book = client.get(f"/api/v2/books/{resp.json()['id']}").json()
    assert [(b["source"][0]["block_kind"], b["target"][0]["block_kind"]) for b in book["beads"]] == [
        ("heading", "heading"),
        ("paragraph", "paragraph"),
    ]


def test_title_defaults_to_source_stem(client):
    assert _import(client).json()["title"] == "contrefeu"


def test_latin1_warning_is_prefixed(client):
    resp = _import(client, target=("sacro.txt", "Marie arrivò.".encode("latin-1")))
    assert resp.json()["warnings"] == ["target: The text file is not valid UTF-8; it was read as Latin-1."]


def test_list_only_books_with_counts(client):
    client.post("/api/v1/projects", json={"title": "Old", "source_lang": "fr", "target_lang": "it"})
    book_id = _import(client).json()["id"]
    assert client.get("/api/v2/books").json() == [{
        "id": book_id, "title": "contrefeu", "source_lang": "fr", "target_lang": "it",
        "bead_count": 3, "reviewed_count": 0,
    }]


def test_excluded_block_sits_after_previous_bead(client, db_path):
    book_id = _import(client).json()["id"]
    book = client.get(f"/api/v2/books/{book_id}").json()
    first_bead = book["beads"][0]["id"]
    block_id = book["beads"][1]["source"][0]["block_id"]  # "Il pleuvait."
    conn = get_connection(db_path)
    try:
        with transaction(conn):
            exclude_block(conn, book_id, block_id)
    finally:
        conn.close()

    book = client.get(f"/api/v2/books/{book_id}").json()
    assert len(book["beads"]) == 3
    assert book["excluded"] == [{
        "block_id": block_id, "side": "source", "kind": "paragraph", "page": None,
        "segments": [{"segment_id": book["excluded"][0]["segments"][0]["segment_id"], "text": "Il pleuvait."}],
        "after_bead_id": first_bead,
    }]


def _pdf(pages) -> bytes:
    """PDF bytes, one portrait page per list of `(x, baseline_y, text, fontsize)`, in an embedded font (as in
    `tests/test_extract_pdf.py`)."""
    pymupdf = pytest.importorskip("pymupdf")
    font = pymupdf.Font("helv")
    doc = pymupdf.open()
    for lines in pages:
        page = doc.new_page(width=439, height=651)
        page.insert_font(fontname="F0", fontbuffer=font.buffer)
        for x, y, text, size in lines:
            page.insert_text((x, y), text, fontsize=size, fontname="F0")
    return doc.tobytes()


def _pdf_book(paragraphs):
    """One indented paragraph per page, and the page number at the bottom."""
    return _pdf([[(74, 120, text, 11.9), (210, 576, str(k + 1), 10)] for k, text in enumerate(paragraphs)])


def test_import_pdf_excludes_page_numbers(client):
    resp = _import(
        client,
        source=("livre.pdf", _pdf_book(["Marie arriva.", "Il pleuvait.", "Paul partit."])),
        target=("libro.pdf", _pdf_book(["Marie arrivò.", "Pioveva.", "Paul partì."])),
    )
    assert resp.status_code == 201
    book = client.get(f"/api/v2/books/{resp.json()['id']}").json()
    assert [(_texts(b, "source"), _texts(b, "target")) for b in book["beads"]] == [
        (["Marie arriva."], ["Marie arrivò."]),
        (["Il pleuvait."], ["Pioveva."]),
        (["Paul partit."], ["Paul partì."]),
    ]
    assert book["beads"][0]["source"][0]["block_kind"] == "paragraph"
    assert sorted((b["side"], b["kind"], b["page"], b["segments"][0]["text"]) for b in book["excluded"]) == sorted(
        (side, "page_number", k, str(k)) for side in ("source", "target") for k in (1, 2, 3)
    )


def test_unreadable_pdf_is_refused_and_writes_nothing(client):
    resp = _import(client, source=("livre.pdf", b"%PDF-1.4"))
    assert resp.status_code == 400
    assert resp.json()["detail"] == "The file is not a readable PDF"
    assert client.get("/api/v2/books").json() == []


def test_empty_txt_is_refused_and_writes_nothing(client):
    resp = _import(client, target=("vide.txt", b"  \n\n"))
    assert resp.status_code == 400
    assert resp.json()["detail"] == "no text found in vide.txt"
    assert client.get("/api/v2/books").json() == []


def test_unknown_extension_is_refused(client):
    assert _import(client, source=("livre.epub", b"x")).status_code == 400


def test_unknown_book_is_404(client):
    assert client.get("/api/v2/books/nope").status_code == 404
    assert client.get("/api/v2/books/nope/check").status_code == 404


def test_old_project_is_not_a_book(client):
    project_id = client.post(
        "/api/v1/projects", json={"title": "Old", "source_lang": "fr", "target_lang": "it"}
    ).json()["id"]
    assert client.get(f"/api/v2/books/{project_id}").status_code == 404


def test_check_is_clean(client):
    book_id = _import(client).json()["id"]
    assert client.get(f"/api/v2/books/{book_id}/check").json() == []
