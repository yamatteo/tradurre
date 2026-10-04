"""Tests for the front/back matter rule (tradurre/services/matter.py), on synthetic blocks only."""

from io import BytesIO

from docx import Document

from tradurre.services.extract import ExtractedBlock, extract
from tradurre.services.matter import mark_matter

LONG = " ".join(["parole"] * 70)  # a long paragraph: body text


def _p(text: str, page: int | None = None) -> ExtractedBlock:
    return ExtractedBlock("paragraph", text, page)


def _h(text: str, page: int | None = None) -> ExtractedBlock:
    return ExtractedBlock("heading", text, page)


def _kinds(blocks: list[ExtractedBlock]) -> list[str]:
    return [b.kind for b in mark_matter(blocks)]


def test_marked_front_run_before_chapter_one():
    blocks = [_p("Author", 1), _h("The Title", 2), _p("Du même auteur", 2), _p("A novel", 2),
              _h("1", 3), _p("short opening line", 3), _p(LONG, 3)]
    assert _kinds(blocks) == ["front_matter", "heading", "front_matter", "front_matter",
                              "heading", "paragraph", "paragraph"]


def test_no_marker_no_exclusion():
    blocks = [_p("Author", 1), _p("A novel", 2), _h("1", 3), _p(LONG, 3), _p("The end", 4), _p("Printer", 5)]
    assert _kinds(blocks) == [b.kind for b in blocks]


def test_front_region_takes_the_marker_page_only():
    blocks = [_p("© 2020", 1), _p("same page as the marker", 1), _p("an epigraph", 2), _h("I", 3), _p(LONG, 3)]
    assert _kinds(blocks) == ["front_matter", "front_matter", "paragraph", "heading", "paragraph"]


def test_long_paragraph_starts_the_body_without_chapter_numbers():
    blocks = [_p("ISBN 978", 1), _p(LONG, 2), _p("short dialogue", 2), _p(LONG, 3)]
    assert _kinds(blocks) == ["front_matter", "paragraph", "paragraph", "paragraph"]


def test_short_last_paragraph_survives_before_the_colophon_page():
    blocks = [_h("1", 1), _p(LONG, 1), _p(LONG, 9), _p("The last short paragraph.", 10),
              _p("unmarked colophon line", 11), _p("Finito di stampare", 11), _p("Printer", 11)]
    assert _kinds(blocks) == ["heading", "paragraph", "paragraph", "paragraph",
                              "back_matter", "back_matter", "back_matter"]


def test_marker_inside_the_body_changes_nothing():
    blocks = [_h("1", 1), _p(LONG, 1), _p("Les Éditions du coin", 2), _p(LONG, 3), _p("The end.", 3)]
    assert _kinds(blocks) == [b.kind for b in blocks]


def test_without_pages_regions_go_by_index():
    blocks = [_p("Titolo originale"), _p("an epigraph"), _p(LONG), _p("Last."), _p("Editore"), _p("Printer")]
    assert _kinds(blocks) == ["front_matter", "paragraph", "paragraph", "paragraph", "back_matter", "back_matter"]


def test_other_kinds_are_never_rekinded():
    blocks = [ExtractedBlock("footnote", "© 2020", 1), _p("x", 1), _h("1", 2), _p(LONG, 2),
              ExtractedBlock("page_number", "ISBN", 3), _h("Notes", 3)]
    assert _kinds(blocks) == ["footnote", "front_matter", "heading", "paragraph", "page_number", "heading"]


def test_empty_and_all_short_unmarked():
    assert mark_matter([]) == []
    blocks = [_p("one"), _p("two"), _h("Three")]
    assert _kinds(blocks) == ["paragraph", "paragraph", "heading"]


def test_input_is_not_mutated():
    blocks = [_p("ISBN 978", 1), _p(LONG, 2)]
    mark_matter(blocks)
    assert blocks[0].kind == "paragraph"


def test_extract_applies_the_rule_to_docx():
    doc = Document()
    doc.add_paragraph("© Une maison d'édition")
    doc.add_paragraph(LONG)
    buffer = BytesIO()
    doc.save(buffer)
    assert [b.kind for b in extract("livre.docx", buffer.getvalue()).blocks] == ["front_matter", "paragraph"]
