"""Tests for tradurre.services.doc_adapter: txt/docx/pdf -> shared text extraction."""

from io import BytesIO

import pytest
from docx import Document

from tradurre.services.doc_adapter import load_as_text, normalize_ocr_artifacts


class TestLoadAsTextTxt:
    def test_decodes_utf8(self):
        text = load_as_text("book.txt", "Héllo wörld.".encode("utf-8"))
        assert text == "Héllo wörld."

    def test_falls_back_to_latin1_on_bad_utf8(self):
        content = "café".encode("latin-1")
        text = load_as_text("book.txt", content)
        assert "caf" in text


class TestLoadAsTextDocx:
    def test_extracts_paragraphs_joined_by_blank_lines(self):
        doc = Document()
        doc.add_paragraph("First paragraph.")
        doc.add_paragraph("Second paragraph.")
        buf = BytesIO()
        doc.save(buf)

        text = load_as_text("book.docx", buf.getvalue())
        assert text == "First paragraph.\n\nSecond paragraph."

    def test_skips_empty_paragraphs(self):
        doc = Document()
        doc.add_paragraph("First.")
        doc.add_paragraph("")
        doc.add_paragraph("Second.")
        buf = BytesIO()
        doc.save(buf)

        text = load_as_text("book.docx", buf.getvalue())
        assert text == "First.\n\nSecond."


class TestLoadAsTextPdf:
    def test_extracts_text_from_pdf(self):
        fitz = pytest.importorskip("fitz")

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Hello from the PDF.")
        buf = doc.tobytes()
        doc.close()

        text = load_as_text("book.pdf", buf)
        assert "Hello from the PDF." in text

    def test_missing_pymupdf_raises_clear_error(self, monkeypatch):
        import builtins

        real_import = builtins.__import__

        def fake_import(name, *args, **kwargs):
            if name == "fitz":
                raise ImportError("no fitz")
            return real_import(name, *args, **kwargs)

        monkeypatch.setattr(builtins, "__import__", fake_import)
        with pytest.raises(RuntimeError, match="pdf"):
            load_as_text("book.pdf", b"%PDF-1.4 fake")


class TestUnsupportedExtension:
    def test_raises_value_error(self):
        with pytest.raises(ValueError):
            load_as_text("book.rtf", b"content")


class TestNormalizeOcrArtifacts:
    def test_fixes_ligatures(self):
        assert normalize_ocr_artifacts("difﬁcult") == "difficult"

    def test_collapses_horizontal_whitespace_only(self):
        text = normalize_ocr_artifacts("a  b\tc\nd")
        assert text == "a b c\nd"

    def test_strips_control_chars(self):
        text = normalize_ocr_artifacts("a\x00b\x0bc")
        assert text == "abc"

    def test_does_not_alter_words(self):
        original = "Marie-Ange went to Pontorgueil."
        assert normalize_ocr_artifacts(original) == original
