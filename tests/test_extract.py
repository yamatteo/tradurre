"""Tests for structured extraction of .txt and .docx files (tradurre.services.extract)."""

from io import BytesIO

import pytest
from docx import Document

from tradurre.services.extract import EXCLUDED_KINDS, ExtractedBlock, extract


def test_txt_with_blank_lines_joins_wrapped_lines():
    text = "Le soir tombait\nsur la ville.\n\n\nIl partit   sans\tun mot.\n\nFin.\n"
    result = extract("livre.txt", text.encode("utf-8"))
    assert result.blocks == [
        ExtractedBlock("paragraph", "Le soir tombait sur la ville."),
        ExtractedBlock("paragraph", "Il partit sans un mot."),
        ExtractedBlock("paragraph", "Fin."),
    ]
    assert result.warnings == []


def test_txt_without_blank_lines_gives_one_block_per_line():
    text = "  Première ligne.\r\nDeuxième  ligne.\r\nTroisième ligne."
    result = extract("livre.TXT", text.encode("utf-8"))
    assert [b.text for b in result.blocks] == ["Première ligne.", "Deuxième ligne.", "Troisième ligne."]
    assert {b.kind for b in result.blocks} == {"paragraph"}


def test_txt_cleans_artifacts():
    result = extract("livre.txt", "Un ﬁlm\x01 muet.".encode("utf-8"))
    assert [b.text for b in result.blocks] == ["Un film muet."]


def test_txt_latin1_fallback_warns():
    result = extract("livre.txt", "L'été était chaud.".encode("latin-1"))
    assert [b.text for b in result.blocks] == ["L'été était chaud."]
    assert len(result.warnings) == 1
    assert "UTF-8" in result.warnings[0]


def test_docx_heading_and_paragraphs():
    doc = Document()
    doc.add_heading("Chapitre premier", level=1)
    doc.add_paragraph("Le soir tombait.")
    doc.add_paragraph("   ")
    doc.add_paragraph("Il partit sans un mot.")
    doc.add_heading("Contrefeu", level=0)  # style "Title"
    buffer = BytesIO()
    doc.save(buffer)
    result = extract("livre.docx", buffer.getvalue())
    assert result.blocks == [
        ExtractedBlock("heading", "Chapitre premier"),
        ExtractedBlock("paragraph", "Le soir tombait."),
        ExtractedBlock("paragraph", "Il partit sans un mot."),
        ExtractedBlock("heading", "Contrefeu"),
    ]
    assert result.warnings == []


def test_pdf_not_implemented_yet():
    with pytest.raises(NotImplementedError):
        extract("livre.pdf", b"%PDF-1.4")


def test_unknown_extension():
    with pytest.raises(ValueError):
        extract("livre.epub", b"")


def test_excluded_kinds():
    assert EXCLUDED_KINDS == {"running_head", "page_number", "footnote", "front_matter", "back_matter"}
